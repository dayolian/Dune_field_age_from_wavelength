#!/usr/bin/env python3
"""
measure_wavelengths_nearest.py

Production wavelength extraction using the 'nearest_tile' strategy.

For each tile:
  - run the FFT, extract all prominent radial peaks (with orientation/azimuth)
  - choose the peak nearest the hand-measured wavelength for that tile
  - if the nearest peak is within TOL, mark high-confidence;
    otherwise still output it, marked low-confidence
  - record ALL candidate peaks so every choice is auditable

NOTE ON METHOD: this uses the hand-measured lambda to select which FFT peak is
the dune. The output is therefore an FFT-REFINED wavelength guided by the field
estimate, not an independent prediction. It is a validation/refinement tool.

Output CSV columns (one row per tile):
  region, lat, lon, tile_basename,
  hand_lambda_m,               # the field estimate used as target
  chosen_lambda_m,             # FFT peak nearest the hand value
  chosen_crest_az_deg,         # crest azimuth of the chosen peak
  chosen_prom_frac, chosen_rel_power,
  within_tol, confidence,      # high/low
  pct_diff_from_hand,          # |chosen - hand| / hand * 100
  n_candidates,
  all_candidates               # "lambda@az(prom);lambda@az(prom);..." for audit

Usage:
  python measure_wavelengths_nearest.py \
      --tiles_dir tiles \
      --truth dune_migration_and_lambda.xlsx \
      --out results/final_wavelengths.csv \
      --tol 0.25
"""
import os, csv, argparse
import numpy as np
import pandas as pd
import openpyxl
from scipy.signal import find_peaks, peak_prominences

from fft_full_geotiff_multimode import (
    load_full_image, high_pass_filter, apply_hann_window, compute_fft_power,
    polar_power_histogram_logk, radial_spectrum_directional,
    theta_to_crest_azimuth, sigma_px_from_cutoff,
    MIN_LAMBDA_M, MAX_LAMBDA_M_GLOBAL, N_ANGLE_BINS, N_RADIAL_BINS,
)

HP_CUTOFF = 3000.0
PROM_FRAC_MIN = 0.15


def tile_name(region, lat, lon):
    r = region.replace(" ", "_")
    return f"{r}_lat{lat:+05.2f}_lon{lon:+06.2f}".replace(".", "p").replace("+", "")


def load_truth(xlsx):
    wb = openpyxl.load_workbook(xlsx); ws = wb["Sheet1"]
    hdr = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col = {n: i + 1 for i, n in enumerate(hdr) if n}
    rows = []
    for r in range(2, ws.max_row + 1):
        lat = ws.cell(row=r, column=col["lat"]).value
        if lat is None:
            continue
        fill = ws.cell(row=r, column=col["FFT lambda"]).fill
        tag = fill.start_color.rgb if fill and fill.patternType else "NONE"
        colour = "yellow" if tag == "FFFFFF00" else ("green" if tag == "FF92D050" else "none")
        rows.append(dict(
            region=str(ws.cell(row=r, column=col["region"]).value),
            lat=float(lat), lon=float(ws.cell(row=r, column=col["lon"]).value),
            hand=ws.cell(row=r, column=col["actual lambda (m)"]).value,
            colour=colour))
    df = pd.DataFrame(rows)
    df["hand"] = pd.to_numeric(df["hand"], errors="coerce")
    return df


def extract_peaks(tif):
    """All prominent radial peaks with orientation/azimuth, using directional spectrum."""
    img, pix = load_full_image(tif)
    ny, nx = img.shape
    lam_max = min(MAX_LAMBDA_M_GLOBAL, 0.7 * min(ny, nx) * pix)
    sigma = sigma_px_from_cutoff(HP_CUTOFF, pix)
    win = apply_hann_window(high_pass_filter(img, sigma))
    P, kx, ky = compute_fft_power(win, pix)
    tc, kc, H = polar_power_histogram_logk(
        P, kx, ky, n_angle_bins=N_ANGLE_BINS, n_radial_bins=N_RADIAL_BINS,
        min_lambda_m=MIN_LAMBDA_M, max_lambda_m=lam_max)
    if H is None:
        return []
    S_k, _ = radial_spectrum_directional(tc, kc, H)
    idx, _ = find_peaks(S_k)
    if len(idx) == 0:
        return []
    prom, _, _ = peak_prominences(S_k, idx)
    heights = S_k[idx]
    pf = prom / (heights + 1e-8)
    main = heights.max()
    peaks = []
    for j, ik in enumerate(idx):
        if kc[ik] <= 0:
            continue
        lam = 1.0 / kc[ik]
        if lam < MIN_LAMBDA_M or lam > lam_max or pf[j] < PROM_FRAC_MIN:
            continue
        # crest azimuth: orientation bin with most power at this radial bin
        col = H[:, ik]
        ith = int(np.argmax(col)) if col.max() > 0 else 0
        _, crest_az = theta_to_crest_azimuth(np.degrees(tc[ith]))
        peaks.append(dict(
            lambda_m=lam, crest_az=crest_az,
            power=float(heights[j]), rel_power=float(heights[j] / (main + 1e-8)),
            prom_frac=float(pf[j])))
    peaks.sort(key=lambda p: p["lambda_m"])
    return peaks


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--tiles_dir", default="tiles")
    ap.add_argument("--truth", default="dune_migration_and_lambda.xlsx")
    ap.add_argument("--out", default="results/final_wavelengths.csv")
    ap.add_argument("--tol", type=float, default=0.25)
    args = ap.parse_args()

    truth = load_truth(args.truth)
    # dune rows only; green (no-dune) tiles are excluded from measurement
    dune = truth[truth["colour"].isin(["yellow", "none"])].copy()

    os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
    fields = ["region", "lat", "lon", "tile_basename", "hand_lambda_m",
              "chosen_lambda_m", "chosen_crest_az_deg", "chosen_prom_frac",
              "chosen_rel_power", "within_tol", "confidence",
              "pct_diff_from_hand", "n_candidates", "all_candidates"]

    n_written = n_high = n_low = n_missing = 0
    with open(args.out, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for _, row in dune.iterrows():
            base = tile_name(row["region"], row["lat"], row["lon"])
            tif = os.path.join(args.tiles_dir, base + ".tif")
            hand = row["hand"]
            if not os.path.exists(tif) or hand is None or not np.isfinite(hand):
                n_missing += 1
                continue
            peaks = extract_peaks(tif)
            if not peaks:
                n_missing += 1
                continue
            # choose peak nearest the hand value
            chosen = min(peaks, key=lambda p: abs(p["lambda_m"] - hand))
            pct = abs(chosen["lambda_m"] - hand) / hand * 100
            within = pct <= args.tol * 100
            conf = "high" if within else "low"
            if within: n_high += 1
            else:      n_low += 1
            cand_str = ";".join(
                f"{p['lambda_m']:.0f}@{p['crest_az']:.0f}(p{p['prom_frac']:.2f})"
                for p in peaks)
            w.writerow(dict(
                region=row["region"], lat=row["lat"], lon=row["lon"],
                tile_basename=base,
                hand_lambda_m=round(hand, 1),
                chosen_lambda_m=round(chosen["lambda_m"], 1),
                chosen_crest_az_deg=round(chosen["crest_az"], 1),
                chosen_prom_frac=round(chosen["prom_frac"], 3),
                chosen_rel_power=round(chosen["rel_power"], 3),
                within_tol=within, confidence=conf,
                pct_diff_from_hand=round(pct, 1),
                n_candidates=len(peaks),
                all_candidates=cand_str))
            n_written += 1

    print(f"Wrote {n_written} tiles to {args.out}")
    print(f"  high-confidence (within {args.tol:.0%}): {n_high}  ({n_high/max(n_written,1)*100:.1f}%)")
    print(f"  low-confidence  (nearest but > tol):   {n_low}  ({n_low/max(n_written,1)*100:.1f}%)")
    print(f"  tiles missing/no-peak (not written):   {n_missing}")


if __name__ == "__main__":
    main()
